import os
import re
from datetime import datetime

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUTS_DIR = os.path.join(os.path.dirname(__file__), "..", "outputs")


# ─────────────────────────────────────────────
# MAIN EXPORT FUNCTION
# ─────────────────────────────────────────────
def export_documents(mom_json: dict, session_id: str) -> tuple:

    os.makedirs(OUTPUTS_DIR, exist_ok=True)

    raw_title = (
        mom_json.get("session_title")
        or mom_json.get("title")
        or "MoM_Report"
    )

    clean_name = re.sub(r"[^\w\s-]", "", str(raw_title))
    clean_name = re.sub(r"\s+", "_", clean_name.strip())
    clean_name = clean_name[:60]

    filename = f"{clean_name}_{session_id[:8]}"

    docx_path = os.path.join(
        OUTPUTS_DIR,
        f"{filename}.docx"
    )

    doc_type = mom_json.get("document_type", "mom")
    if doc_type == "class_notes" or doc_type == "online_session":
        _generate_class_notes_docx(mom_json, docx_path)
    else:
        _generate_docx(mom_json, docx_path)

    return None, f"/outputs/{filename}.docx"


# ─────────────────────────────────────────────
# DOCX GENERATION
# ─────────────────────────────────────────────
def _generate_docx(data: dict, path: str):

    doc = Document()

    # PAGE SETTINGS
    section = doc.sections[0]
    section.top_margin = Pt(45)
    section.bottom_margin = Pt(45)
    section.left_margin = Pt(50)
    section.right_margin = Pt(50)

    # FOOTER
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer_para)

    # ─────────────────────────────────────────
    # 1. MAIN TITLE
    # ─────────────────────────────────────────
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("Minutes of the Meeting")
    run.bold = True
    run.font.size = Pt(16)
    title_para.paragraph_format.space_after = Pt(14)

    # ─────────────────────────────────────────
    # 2. METADATA TABLE (2 COLUMNS)
    # ─────────────────────────────────────────
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.style = "Table Grid"

    meeting_title  = data.get("session_title") or "Meeting Review"
    meeting_no     = data.get("meeting_no") or f"{datetime.now().strftime('%Y-%m')}/01"
    meeting_date   = data.get("date") or datetime.now().strftime("%d.%m.%Y")
    meeting_time   = data.get("time") or "Scheduled Session"
    venue_platform = data.get("venue_platform") or "Google Meet"

    meta_rows = [
        ("Meeting Title",       meeting_title),
        ("Meeting No.",         meeting_no),
        ("Date",                meeting_date),
        ("Time",                meeting_time),
        ("Venue / Platform",    venue_platform),
    ]

    for idx, (label, val) in enumerate(meta_rows):
        row_cells = meta_table.rows[idx].cells
        
        # Label cell
        row_cells[0].text = str(label)
        p0 = row_cells[0].paragraphs[0]
        r0 = p0.runs[0]
        r0.bold = True
        r0.font.size = Pt(10)
        _set_cell_background(row_cells[0], "F2F2F2")

        # Value cell
        row_cells[1].text = str(val)
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.runs[0]
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ─────────────────────────────────────────
    # 3. MEMBERS PRESENT
    # ─────────────────────────────────────────
    mem_header = doc.add_paragraph()
    r_mem = mem_header.add_run("Members Present:")
    r_mem.bold = True
    r_mem.font.size = Pt(11)
    mem_header.paragraph_format.space_after = Pt(4)

    members = data.get("members_present") or data.get("participants") or ["Attendees"]
    if isinstance(members, str):
        members = [members]

    for member in members:
        p_mem = doc.add_paragraph(style="List Bullet")
        r_m = p_mem.add_run(str(member))
        r_m.font.size = Pt(10)
        p_mem.paragraph_format.space_after = Pt(2)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(8)

    # ─────────────────────────────────────────
    # 4. POINTS DISCUSSED
    # ─────────────────────────────────────────
    disc_header = doc.add_paragraph()
    r_disc = disc_header.add_run("Points Discussed")
    r_disc.bold = True
    r_disc.font.size = Pt(13)
    disc_header.paragraph_format.space_after = Pt(8)

    points_discussed = data.get("points_discussed") or []
    
    # Auto-convert legacy/fallback structures
    if not points_discussed:
        categories = data.get("categories") or []
        if categories:
            for cat in categories:
                points_discussed.append({
                    "category_name": cat.get("name") or "General Discussion",
                    "points": cat.get("points") or ["Discussion conducted."]
                })
        else:
            points_discussed = [{
                "category_name": "General Discussion",
                "points": ["The meeting proceedings were conducted as per agenda."]
            }]

    for cat_item in points_discussed:
        if not isinstance(cat_item, dict):
            continue

        cat_name = cat_item.get("category_name") or "Discussion"
        points   = cat_item.get("points") or []
        if isinstance(points, str):
            points = [points]

        # Category Header
        p_cat = doc.add_paragraph()
        r_cname = p_cat.add_run(f"Category: {cat_name}")
        r_cname.bold = True
        r_cname.font.size = Pt(10.5)
        p_cat.paragraph_format.space_before = Pt(4)
        p_cat.paragraph_format.space_after = Pt(3)

        for pt in points:
            p_pt = doc.add_paragraph(style="List Bullet")
            r_pt = p_pt.add_run(str(pt))
            r_pt.font.size = Pt(10)
            p_pt.paragraph_format.space_after = Pt(2)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ─────────────────────────────────────────
    # 5. RESPONSIBILITY & TARGET DATE
    # ─────────────────────────────────────────
    resp_header = doc.add_paragraph()
    r_resp = resp_header.add_run("Responsibility & Target Date")
    r_resp.bold = True
    r_resp.font.size = Pt(13)
    resp_header.paragraph_format.space_after = Pt(8)

    resp_matrix = data.get("responsibility_matrix") or []
    if not resp_matrix:
        for cat_item in points_discussed:
            c_name = cat_item.get("category_name") if isinstance(cat_item, dict) else "Discussion"
            resp_matrix.append({
                "category_name": c_name,
                "responsibility": "All Members",
                "target_date": "Continuous"
            })

    resp_table = doc.add_table(rows=1, cols=3)
    resp_table.style = "Table Grid"

    table_headers = ["Category", "Responsibility", "Target Date"]
    hdr_cells = resp_table.rows[0].cells
    for i, h in enumerate(table_headers):
        hdr_cells[i].text = h
        p_h = hdr_cells[i].paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_h = p_h.runs[0]
        r_h.bold = True
        r_h.font.size = Pt(10)
        _set_cell_background(hdr_cells[i], "D9D9D9")

    for item in resp_matrix:
        if not isinstance(item, dict):
            continue

        row_cells = resp_table.add_row().cells
        
        row_cells[0].text = str(item.get("category_name") or "General")
        row_cells[1].text = str(item.get("responsibility") or "All")
        row_cells[2].text = str(item.get("target_date") or "Continuous")

        for c_idx in range(3):
            p_c = row_cells[c_idx].paragraphs[0]
            if len(p_c.runs) > 0:
                p_c.runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ─────────────────────────────────────────
    # 6. INFORMATION ITEMS
    # ─────────────────────────────────────────
    info_header = doc.add_paragraph()
    r_info = info_header.add_run("Information Items")
    r_info.bold = True
    r_info.font.size = Pt(13)
    info_header.paragraph_format.space_after = Pt(8)

    info_items = data.get("information_items") or []
    if isinstance(info_items, str):
        info_items = [info_items]

    if not info_items:
        info_items = [
            "All members are requested to review the notes and complete assigned tasks.",
            "Schedule for the next review session will be communicated shortly."
        ]

    for idx, item in enumerate(info_items, start=1):
        p_item = doc.add_paragraph()
        r_num = p_item.add_run(f"{idx}. ")
        r_num.bold = True
        r_num.font.size = Pt(10)
        r_txt = p_item.add_run(str(item))
        r_txt.font.size = Pt(10)
        p_item.paragraph_format.space_after = Pt(3)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ─────────────────────────────────────────
    # 7. DISTRIBUTION & SIGN-OFF
    # ─────────────────────────────────────────
    copy_to = data.get("copy_to") or ["All Members"]
    if isinstance(copy_to, str):
        copy_to = [copy_to]

    p_ct = doc.add_paragraph()
    r_ct = p_ct.add_run("Copy To:")
    r_ct.bold = True
    r_ct.font.size = Pt(10.5)
    p_ct.paragraph_format.space_after = Pt(3)

    for item in copy_to:
        p_c = doc.add_paragraph(style="List Bullet")
        r_c = p_c.add_run(str(item))
        r_c.font.size = Pt(10)
        p_c.paragraph_format.space_after = Pt(2)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    copy_sub = data.get("copy_submitted_to") or ["Management"]
    if isinstance(copy_sub, str):
        copy_sub = [copy_sub]

    p_csub = doc.add_paragraph()
    r_csub = p_csub.add_run("Copy Submitted To:")
    r_csub.bold = True
    r_csub.font.size = Pt(10.5)
    p_csub.paragraph_format.space_after = Pt(3)

    for item in copy_sub:
        p_cs = doc.add_paragraph(style="List Bullet")
        r_cs = p_cs.add_run(str(item))
        r_cs.font.size = Pt(10)
        p_cs.paragraph_format.space_after = Pt(2)

    # SIGNATURE BLOCK
    p_sign_space = doc.add_paragraph()
    p_sign_space.paragraph_format.space_before = Pt(16)
    p_sign_space.paragraph_format.space_after = Pt(2)
    
    r_line = p_sign_space.add_run("_______________________________")
    r_line.bold = True

    sig_name = data.get("signatory_name") or "Meeting Secretary"
    sig_desig = data.get("signatory_designation") or "Convener"
    sig_date = data.get("signature_date") or data.get("date") or datetime.now().strftime("%d.%m.%Y")

    p_sig1 = doc.add_paragraph()
    r_s1 = p_sig1.add_run(str(sig_name))
    r_s1.bold = True
    r_s1.font.size = Pt(10.5)
    p_sig1.paragraph_format.space_after = Pt(2)

    p_sig2 = doc.add_paragraph()
    r_s2 = p_sig2.add_run(str(sig_desig))
    r_s2.italic = True
    r_s2.font.size = Pt(10)
    p_sig2.paragraph_format.space_after = Pt(2)

    p_sig3 = doc.add_paragraph()
    r_s3 = p_sig3.add_run(f"Date: {sig_date}")
    r_s3.bold = True
    r_s3.font.size = Pt(10)

    # SAVE DOCUMENT
    try:
        doc.save(path)
        print("[OK] Standard MoM DOCX saved:", path)
    except Exception as e:
        print("[ERROR] DOCX save failed:", e)
        raise




# ─────────────────────────────────────────────
# CELL BACKGROUND
# ─────────────────────────────────────────────
def _set_cell_background(cell, color):

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")

    shd.set(qn("w:fill"), color)

    tc_pr.append(shd)


# ─────────────────────────────────────────────
# PAGE NUMBER
# ─────────────────────────────────────────────
def _add_page_number(paragraph):

    paragraph.add_run("Page ")

    # PAGE
    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(
        qn("w:fldCharType"),
        "begin"
    )

    run._r.append(fld_begin)

    run = paragraph.add_run()

    instr = OxmlElement("w:instrText")
    instr.set(
        qn("xml:space"),
        "preserve"
    )

    instr.text = "PAGE"

    run._r.append(instr)

    run = paragraph.add_run()

    fld_end = OxmlElement("w:fldChar")

    fld_end.set(
        qn("w:fldCharType"),
        "end"
    )

    run._r.append(fld_end)

    paragraph.add_run(" of ")

    # NUMPAGES
    run = paragraph.add_run()

    fld_begin2 = OxmlElement("w:fldChar")

    fld_begin2.set(
        qn("w:fldCharType"),
        "begin"
    )

    run._r.append(fld_begin2)

    run = paragraph.add_run()

    instr2 = OxmlElement("w:instrText")

    instr2.set(
        qn("xml:space"),
        "preserve"
    )

    instr2.text = "NUMPAGES"

    run._r.append(instr2)

    run = paragraph.add_run()

    fld_end2 = OxmlElement("w:fldChar")

    fld_end2.set(
        qn("w:fldCharType"),
        "end"
    )

    run._r.append(fld_end2)

# ─────────────────────────────────────────────
# ─────────────────────────────────────────────
# CLASS / WEBINAR NOTES DOCX GENERATION
# ─────────────────────────────────────────────
def _generate_class_notes_docx(data: dict, path: str):
    doc = Document()

    # PAGE SETTINGS
    section = doc.sections[0]
    section.top_margin = Pt(45)
    section.bottom_margin = Pt(45)
    section.left_margin = Pt(50)
    section.right_margin = Pt(50)

    # FOOTER
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer_para)

    # 1. MAIN TITLE
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    session_title = data.get("session_title") or data.get("title") or "Class / Webinar Notes"
    run = title_para.add_run("CLASS / WEBINAR NOTES")
    run.bold = True
    run.font.size = Pt(16)
    title_para.paragraph_format.space_after = Pt(12)

    # 2. METADATA TABLE
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Table Grid"

    speaker = data.get("speaker_instructor") or data.get("instructor") or "Speaker / Instructor"
    date_val = data.get("date") or datetime.now().strftime("%Y-%m-%d")
    session_type = data.get("session_type") or "Class / Webinar"

    meta_rows = [
        ("Title", str(session_title)),
        ("Date", str(date_val)),
        ("Speaker / Instructor", str(speaker)),
        ("Session Type", str(session_type)),
    ]

    for idx, (label, val) in enumerate(meta_rows):
        row_cells = meta_table.rows[idx].cells
        row_cells[0].text = label
        p0 = row_cells[0].paragraphs[0]
        r0 = p0.runs[0]
        r0.bold = True
        r0.font.size = Pt(10)
        _set_cell_background(row_cells[0], "F2F2F2")

        row_cells[1].text = val
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.runs[0]
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    section_num = 1

    # 1. Overview
    overview = data.get("overview")
    if overview and str(overview).strip():
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Overview")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        p_txt = doc.add_paragraph(str(overview).strip())
        p_txt.paragraph_format.space_after = Pt(8)
        section_num += 1

    # 2. Topics Covered
    topics_covered = data.get("topics_covered") or []
    if isinstance(topics_covered, str):
        topics_covered = [topics_covered]
    if topics_covered:
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Topics Covered")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        for topic in topics_covered:
            if topic and str(topic).strip():
                p_t = doc.add_paragraph(style="List Bullet")
                r_t = p_t.add_run(str(topic).strip())
                r_t.font.size = Pt(10)
                p_t.paragraph_format.space_after = Pt(2)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        section_num += 1

    # 3. Detailed Notes
    detailed_notes = data.get("detailed_notes") or []
    if detailed_notes:
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Detailed Notes")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        for sub_idx, note_item in enumerate(detailed_notes, start=1):
            if isinstance(note_item, dict):
                t_title = note_item.get("topic_title") or note_item.get("topic_name") or f"Topic {sub_idx}"
                explanation = note_item.get("explanation") or note_item.get("summary") or ""
                key_pts = note_item.get("key_points") or []

                p_sub = doc.add_paragraph()
                r_st = p_sub.add_run(f"{section_num - 1}.{sub_idx} {t_title}")
                r_st.bold = True
                r_st.font.size = Pt(11)
                p_sub.paragraph_format.space_before = Pt(6)
                p_sub.paragraph_format.space_after = Pt(3)

                if explanation:
                    p_exp = doc.add_paragraph(str(explanation).strip())
                    p_exp.paragraph_format.space_after = Pt(4)

                if key_pts:
                    if isinstance(key_pts, str):
                        key_pts = [key_pts]
                    for kp in key_pts:
                        p_kp = doc.add_paragraph(style="List Bullet")
                        p_kp.add_run(str(kp).strip())
                        p_kp.paragraph_format.space_after = Pt(2)
            elif isinstance(note_item, str):
                p_txt = doc.add_paragraph(style="List Bullet")
                p_txt.add_run(note_item.strip())
                p_txt.paragraph_format.space_after = Pt(2)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        section_num += 1

    # 4. Important Concepts & Definitions
    important_concepts = data.get("important_concepts") or []
    if important_concepts:
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Important Concepts")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        for concept in important_concepts:
            if isinstance(concept, dict):
                term = concept.get("term_or_concept") or concept.get("term") or ""
                defn = concept.get("definition_or_explanation") or concept.get("explanation") or ""
                p_c = doc.add_paragraph(style="List Bullet")
                if term:
                    r_term = p_c.add_run(f"{term}: ")
                    r_term.bold = True
                p_c.add_run(str(defn).strip())
                p_c.paragraph_format.space_after = Pt(2)
            elif isinstance(concept, str):
                p_c = doc.add_paragraph(style="List Bullet")
                p_c.add_run(concept.strip())
                p_c.paragraph_format.space_after = Pt(2)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        section_num += 1

    # 5. Examples / Demonstrations
    examples = data.get("examples_demonstrations") or data.get("examples") or []
    if isinstance(examples, str):
        examples = [examples]
    if examples:
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Examples / Demonstrations")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        for ex in examples:
            if ex and str(ex).strip():
                p_ex = doc.add_paragraph(style="List Bullet")
                p_ex.add_run(str(ex).strip())
                p_ex.paragraph_format.space_after = Pt(2)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        section_num += 1

    # 6. Code / Technical Examples
    code_examples = data.get("code_technical_examples") or []
    if code_examples:
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Code / Technical Examples")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        for item in code_examples:
            if isinstance(item, dict):
                lang = item.get("language_or_context") or "Code"
                code = item.get("code_snippet") or ""
                expl = item.get("explanation") or ""

                p_l = doc.add_paragraph()
                r_l = p_l.add_run(f"Language / Context: {lang}")
                r_l.bold = True
                p_l.paragraph_format.space_before = Pt(4)

                if code:
                    p_code = doc.add_paragraph()
                    r_code = p_code.add_run(code)
                    r_code.font.name = 'Consolas'
                    r_code.font.size = Pt(9.5)
                    p_code.paragraph_format.left_indent = Pt(15)
                    p_code.paragraph_format.space_after = Pt(4)

                if expl:
                    p_e = doc.add_paragraph(f"Explanation: {expl}")
                    p_e.paragraph_format.space_after = Pt(4)
            elif isinstance(item, str):
                p_c = doc.add_paragraph(item)
                p_c.paragraph_format.space_after = Pt(4)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        section_num += 1

    # 7. Questions & Answers
    qna = data.get("questions_and_answers") or data.get("doubts_and_clarifications") or []
    if qna:
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Questions & Answers")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        for item in qna:
            if isinstance(item, dict):
                q = item.get("question") or ""
                a = item.get("answer") or ""
                p_qa = doc.add_paragraph()
                r_q = p_qa.add_run(f"Q: {q}\n")
                r_q.bold = True
                r_a = p_qa.add_run(f"A: {a}")
                p_qa.paragraph_format.space_after = Pt(4)
            elif isinstance(item, str):
                p_qa = doc.add_paragraph(item)
                p_qa.paragraph_format.space_after = Pt(4)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        section_num += 1

    # 8. Practical Tips & Recommendations
    tips = data.get("practical_tips") or []
    if isinstance(tips, str):
        tips = [tips]
    if tips:
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Practical Tips & Recommendations")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        for tip in tips:
            if tip and str(tip).strip():
                p_tip = doc.add_paragraph(style="List Bullet")
                p_tip.add_run(str(tip).strip())
                p_tip.paragraph_format.space_after = Pt(2)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        section_num += 1

    # 9. Key Takeaways
    takeaways = data.get("key_takeaways") or []
    if isinstance(takeaways, str):
        takeaways = [takeaways]
    if takeaways:
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Key Takeaways")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        for kw in takeaways:
            if kw and str(kw).strip():
                p_kw = doc.add_paragraph(style="List Bullet")
                p_kw.add_run(str(kw).strip())
                p_kw.paragraph_format.space_after = Pt(2)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        section_num += 1

    # 10. Final Summary
    summary = data.get("final_summary") or data.get("session_summary")
    if summary and str(summary).strip():
        p_hdr = doc.add_paragraph()
        r_hdr = p_hdr.add_run(f"{section_num}. Final Summary")
        r_hdr.bold = True
        r_hdr.font.size = Pt(13)
        p_hdr.paragraph_format.space_before = Pt(10)
        p_hdr.paragraph_format.space_after = Pt(4)

        p_s = doc.add_paragraph(str(summary).strip())
        p_s.paragraph_format.space_after = Pt(6)

    # SAVE DOCUMENT
    try:
        doc.save(path)
        print("[OK] Class Notes DOCX saved:", path)
    except Exception as e:
        print("[ERROR] DOCX save failed:", e)
        raise